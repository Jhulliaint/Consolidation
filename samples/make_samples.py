"""Genere des classeurs "Management accounts" de demonstration.

Le format reproduit celui CONSTATE dans les fichiers Mage reels
(feuilles Cover / Mapping accounts / Group COA / Trial balance,
convention de signe DEBIT (+) / CREDIT (-), lignes de controle en pied).

Ces donnees sont FICTIVES : elles servent uniquement a faire tourner et tester
la chaine complete en l'absence des fichiers de production.
"""
from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).resolve().parent


def _cover(wb: Workbook, entity: str, currency: str, year: int) -> None:
    ws = wb.create_sheet("Cover")
    ws["A1"] = entity
    ws["A3"] = "MANAGEMENT ACCOUNTS"
    ws["A5"] = "12/31"
    ws["C7"], ws["E7"] = "FISCAL YEAR", year
    ws["C9"], ws["E9"] = "LOCAL CURRENCY", currency
    ws["C11"], ws["F11"] = "REPORTING PERIOD START DATE :", date(year, 1, 1)
    ws["C12"], ws["F12"] = "REPORTING PERIOD END DATE :", date(year, 12, 31)


def _mapping(wb: Workbook, rows: list[tuple[str, str]]) -> None:
    ws = wb.create_sheet("Mapping accounts")
    ws.append(["ACCOUNT ID", "LOCAL ACCOUNT ", "GROUP MAGE COA "])
    for local, group in rows:
        ws.append([None, local, group])


def _group_coa(wb: Workbook, bs: list[str], pl: list[str]) -> None:
    ws = wb.create_sheet("Group COA")
    ws.append(["Account description ", None, None])
    for i, cap in enumerate(bs):
        ws.append([cap, "BALANCE SHEET ACCOUNTS " if i == 0 else None, None])
    for i, cap in enumerate(pl):
        ws.append([cap, "PROFIT & LOSS ACCOUNTS " if i == 0 else None, None])


def _trial_balance(wb: Workbook, entity: str, rows: list[tuple[str, float]]) -> None:
    """Balance en convention DEBIT (+) / CREDIT (-)."""
    ws = wb.create_sheet("Trial balance")
    ws["A1"] = entity
    ws["A2"] = "DEBIT (+)/CREDIT (-)"
    ws["A3"], ws["B3"] = "FOR THE PERIOD ENDED", "12/31"
    ws.append([])
    ws.append(["Account description", "Debit", "Credit", "Ending balance"])
    total = 0.0
    for caption, balance in rows:
        debit = balance if balance > 0 else 0.0
        credit = -balance if balance < 0 else 0.0
        ws.append([caption, debit, credit, balance])
        total += balance
    ws.append([None, None, "Control :", round(total, 6)])


# ---------------------------------------------------------------------------
# Jeu de donnees : MAGE SAS (EUR, mere) + MAGE JAPON KK (JPY, filiale)
# Les deux portent des soldes intragroupe reciproques afin d'exercer
# le moteur d'elimination.
# ---------------------------------------------------------------------------
BS_CAPTIONS = [
    "Cash Bank", "Account Receivable", "Account Receivable - subsidiary Mage Japan",
    "Stock shoes", "Stock shoes subsidiaries", "Accounts Payable",
    "Accounts Payable - Intercompany purchase Mage", "Share capital",
    "Retained earnings/losses", "Profit/loss net income",
]
PL_CAPTIONS = [
    "Shoe Sales", "Shoe Sales subsidiaries", "Opening stock shoes Mage SAS",
    "Closing stock shoes Mage SAS", "Purchases shoes from Mage",
    "Wages and salaries", "Rent", "Depreciation", "Corporate income tax",
]

MAGE_SAS_TB = [
    ("Cash Bank", 120_000.00),
    ("Account Receivable", 260_000.00),
    ("Account Receivable - subsidiary Mage Japan", 90_000.00),
    ("Stock shoes", 380_000.00),
    ("Accounts Payable", -540_000.00),
    ("Share capital", -900_000.00),
    # solde d'equilibrage : pertes accumulees (debit)
    ("Retained earnings/losses", 1_422_000.00),
    # resultat : produits en credit (-), charges en debit (+)
    ("Shoe Sales", -1_250_000.00),
    ("Shoe Sales subsidiaries", -300_000.00),
    ("Opening stock shoes Mage SAS", 330_000.00),
    ("Closing stock shoes Mage SAS", -380_000.00),
    ("Wages and salaries", 560_000.00),
    ("Rent", 175_000.00),
    ("Depreciation", 30_000.00),
    ("Corporate income tax", 3_000.00),
]

MAGE_JAPON_TB = [
    ("Cash Bank", 2_300_000.0),
    ("Account Receivable", 15_400_000.0),
    ("Stock shoes subsidiaries", 6_900_000.0),
    # reciproque exacte des 90 000 EUR portes par MAGE SAS (90 000 x 156,33)
    ("Accounts Payable - Intercompany purchase Mage", -14_069_700.0),
    ("Share capital", -20_000_000.0),
    ("Retained earnings/losses", 12_579_700.0),
    ("Shoe Sales", -62_000_000.0),
    ("Purchases shoes from Mage", 46_890_000.0),
    ("Wages and salaries", 8_000_000.0),
    ("Rent", 4_000_000.0),
]

JP_LOCAL_MAP = [
    ("店舗現金/Cash in Store", "Cash in Hand - cash sales"),
    ("普通預金/Ordinary deposit", "Cash Bank"),
    ("売掛金/Trade account receivable", "Account Receivable"),
    ("商品/Inventory", "Stock shoes subsidiaries"),
]


def build(out_dir: Path) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    made: list[Path] = []

    wb = Workbook()
    wb.remove(wb.active)
    _cover(wb, "MAGE SAS", "EUR", 2025)
    _mapping(wb, [(c, c) for c in BS_CAPTIONS + PL_CAPTIONS])
    _group_coa(wb, BS_CAPTIONS, PL_CAPTIONS)
    _trial_balance(wb, "MAGE SAS", MAGE_SAS_TB)
    p = out_dir / "MA - Mage SAS 2025.xlsx"
    wb.save(p)
    made.append(p)

    wb = Workbook()
    wb.remove(wb.active)
    _cover(wb, "MAGE JAPON KK", "JPY", 2025)
    _mapping(wb, JP_LOCAL_MAP + [(c, c) for c in BS_CAPTIONS + PL_CAPTIONS])
    _group_coa(wb, BS_CAPTIONS, PL_CAPTIONS)
    _trial_balance(wb, "MAGE JAPON KK", MAGE_JAPON_TB)
    p = out_dir / "MA - Mage Japon KK 2025.xlsx"
    wb.save(p)
    made.append(p)
    return made


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "data"
    for path in build(target):
        print("ecrit :", path)
