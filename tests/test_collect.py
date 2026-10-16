"""Dossier de collecte alimente par les scenarios Make (chantier C-A)."""
from __future__ import annotations

import shutil
from datetime import date

from openpyxl import Workbook, load_workbook

import make_samples
from mageconso.cli import main
from mageconso.collect import collect, parse_name
from mageconso.pipeline import RunOptions, run


def _copy(src, folder, name):
    folder.mkdir(parents=True, exist_ok=True)
    dst = folder / name
    shutil.copy(src, dst)
    return dst


def _set_period_end(path, end: date):
    wb = load_workbook(path)
    wb["Cover"]["F12"] = end
    wb.save(path)


def _tree(tmp_path):
    """CLOTURE MENSUELLES/2026/{01,02}/... avec une version corrigee."""
    src = make_samples.build(tmp_path / "src")
    sas, jp = src
    jp_v2 = make_samples.build(
        tmp_path / "src_v2",
        jp_tb=make_samples.rebalance(
            [(c, v * 2 if c == "Cash Bank" else v)
             for c, v in make_samples.MAGE_JAPON_TB]))[1]
    root = tmp_path / "CLOTURE MENSUELLES"
    a = _copy(jp, root / "2026" / "01", "20260115-090000 - MAGE_JAPON_KK - Mage Japon report.xlsx")
    b = _copy(jp_v2, root / "2026" / "01",
              "20260120-170500 - MAGE_JAPON_KK - Mage Japon report revised.xlsx")
    # piece jointe qui n'est pas un management accounts (paie)
    payroll = Workbook()
    payroll.active["A1"] = "Payroll report"
    payroll.save(root / "2026" / "01" / "20260119-120000 - MAGE_JAPON_KK - Payroll 202601.xlsx")
    # cloture suivante, deposee le mois d'apres
    c = _copy(jp, root / "2026" / "02", "20260214-080000 - MAGE_JAPON_KK - Mage Japon report.xlsx")
    _set_period_end(c, date(2026, 1, 31))
    for p in (a, b):
        _set_period_end(p, date(2025, 12, 31))
    # fichier depose a la main, sans prefixe
    _copy(sas, root / "2026" / "01", "MA - Mage SAS 2025.xlsx")
    return root, a, b, c


def test_parse_name(tmp_path):
    p = tmp_path / "20260915-093012 - MAGE_JAPON_KK - report Aug 2026 1.xlsx"
    p.write_bytes(b"")
    ts, entity, original = parse_name(p)
    assert ts.isoformat() == "2026-09-15T09:30:12"
    assert entity == "MAGE_JAPON_KK"
    assert original == "report Aug 2026 1.xlsx"
    manual = tmp_path / "MA - Mage SAS 2025.xlsx"
    manual.write_bytes(b"")
    assert parse_name(manual)[1] is None


def test_latest_version_retained_and_replacement_signalled(cfg, tmp_path):
    root, a, b, _ = _tree(tmp_path)
    col = collect(cfg, root, closing="2025-12-31")
    assert {(f.entity, f.path.name) for f in col.retained} == {
        ("MAGE_JAPON_KK", b.name), ("MAGE_SAS", "MA - Mage SAS 2025.xlsx")}
    assert [f.path for f in col.replaced] == [a]
    assert col.replaced[0].replaced_by == b
    codes = [d.code for d in col.diagnostics]
    assert "COLLECT-REPLACED" in codes
    # la paie et la cloture de janvier sont ecartees, pas consolidees
    reasons = {f.path.name: f.reason for f in col.ignored}
    assert any("Payroll" in n and "Trial balance" in r for n, r in reasons.items())
    assert any("autre cloture (2026-01-31)" == r for r in reasons.values())


def test_consolidation_uses_corrected_version(cfg, tmp_path):
    root, a, b, _ = _tree(tmp_path)
    col = collect(cfg, root, closing="2025-12-31")
    out = run(cfg, col.paths, RunOptions(closing="2025-12-31"), workbooks=col.workbooks)
    files = {r.source_file for r in out.result.audit if r.entity == "MAGE_JAPON_KK"}
    assert files == {b.name}


def test_default_is_latest_closing_per_entity(cfg, tmp_path):
    root, _, _, c = _tree(tmp_path)
    col = collect(cfg, root)
    by_entity = {f.entity: f for f in col.retained}
    assert by_entity["MAGE_JAPON_KK"].path == c
    assert by_entity["MAGE_JAPON_KK"].closing == "2026-01-31"
    # Mage SAS est encore au 31/12 : la divergence de clotures est signalee
    assert "COLLECT-CLOSINGS" in [d.code for d in col.diagnostics]


def test_missing_email_entity_signalled(cfg, tmp_path):
    root, *_ = _tree(tmp_path)
    col = collect(cfg, root, closing="2025-12-31")
    missing = [d.entity for d in col.diagnostics if d.code == "COLLECT-MISSING"]
    assert missing == ["MAGE_ASIA_PACIFIC"]


def test_unknown_root(cfg, tmp_path):
    col = collect(cfg, tmp_path / "absent")
    assert col.diagnostics[0].code == "COLLECT-ROOT"
    assert not col.retained


def test_cli_collect_and_consolidate_from_folder(cfg, tmp_path, capsys):
    root, _, b, _ = _tree(tmp_path)
    assert main(["collect", str(root), "--closing", "2025-12-31"]) == 0
    text = capsys.readouterr().out
    assert "RETENU" in text and "REMPLACE" in text and b.name in text

    target = tmp_path / "out.xlsx"
    rc = main(["consolidate", str(root), "--closing", "2025-12-31", "--force",
               "-o", str(target)])
    assert rc == 0 and target.exists()
    text = capsys.readouterr().out
    assert b.name in text
    sheet = load_workbook(target)["Diagnostics"]
    assert any("COLLECT-REPLACED" in str(c.value) for row in sheet.iter_rows() for c in row)


def test_identical_copy_is_duplicate_not_replacement(cfg, tmp_path):
    root, a, b, _ = _tree(tmp_path)
    # OneDrive renomme "... 1.xlsx" la 2e copie d'un meme fichier
    copy = _copy(b, b.parent, b.stem + " 1.xlsx")
    col = collect(cfg, root, closing="2025-12-31")
    jp = {f.path.name: f.status for f in col.files if f.entity == "MAGE_JAPON_KK"
          and f.closing == "2025-12-31"}
    assert jp == {a.name: "replaced", b.name: "retained", copy.name: "duplicate"}
    replaced = [d for d in col.diagnostics if d.code == "COLLECT-REPLACED"]
    assert len(replaced) == 1 and "2 versions" in replaced[0].message
