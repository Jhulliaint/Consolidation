"""Tests des fonctionnalites : pipeline, vues, taux, rapprochement, export,
verification du parametrage."""
from __future__ import annotations

import copy
from decimal import Decimal
from pathlib import Path

import make_samples as ms
import pytest
from openpyxl import Workbook, load_workbook

from mageconso.checks import check_config
from mageconso.config import norm
from mageconso.models import Severity, Statement
from mageconso.pipeline import RunOptions, export, missing_rates, read_sources, run
from mageconso.rates import apply_rates, check_rate, load_rates_file, rate_needs, write_rates_template
from mageconso.reconcile import load_reference

CENT = Decimal("0.01")


@pytest.fixture(scope="module")
def out(cfg, sample_files):
    return run(cfg, sample_files, RunOptions(by_cost_centre=True))


# ------------------------------------------------------------------- pipeline
def test_pipeline_produces_all_views(out):
    assert {"Consolidated BS", "Consolidated PL", "Detailed BS", "Detailed PL",
            "BS by entity", "PL by entity", "BS by cost center", "PL by cost center",
            "Gross margin"} <= set(out.reports)


def test_entity_view_columns_and_totals(out):
    rep = out.reports["BS by entity"]
    assert rep.columns == ["MAGE_SAS", "MAGE_JAPON_KK", "ELIMINATION", "TOTAL"]
    assert rep.label("MAGE_JAPON_KK") == "MAGE JAPON KK"
    row = next(r for r in rep.rows if r.caption == "Total assets")
    assert row.value("TOTAL") == sum(
        (row.value(c) for c in rep.columns[:-1]), Decimal(0))
    assert abs(row.value("TOTAL") - out.bs.total("Total assets")) < CENT


def test_entity_view_shows_elimination_column(out):
    rep = out.reports["PL by entity"]
    row = next(r for r in rep.rows if r.caption == "Shoe Sales subsidiaries")
    assert row.value("MAGE_SAS") == Decimal(300_000)
    assert row.value("ELIMINATION") == Decimal(-300_000)
    assert row.value("TOTAL") == 0


def test_percent_of_revenue(out):
    pl = out.pl
    rev = next(r for r in pl.rows if r.caption == "Revenue")
    net = next(r for r in pl.rows if r.caption == "Total Net/(loss) Profit")
    assert pl.percent(rev, "TOTAL") == 1
    assert abs(pl.percent(net, "TOTAL") - net.value() / rev.value()) < Decimal("1e-9")


def test_cost_centre_pl_has_cascade(out):
    rep = out.reports["PL by cost center"]
    assert any(r.caption == "Total Net/(loss) Profit" for r in rep.rows)


def test_margin_analysis_matches_reference_rule(out):
    """SHOES = (Shoe Sales + distributeurs) + cout (lignes 52-63 de la
    reference) ; BESPOKE & SHOES = SHOES + BESPOKE."""
    rep = out.reports["Gross margin"]
    rev = next(r for r in rep.rows if r.caption == "Revenue")
    cost = next(r for r in rep.rows if r.caption == "Cost of sales")
    rate = next(r for r in rep.rows if r.caption == "Gross margin %")
    shoes_rev = Decimal(1_250_000) + Decimal(62_000_000) / Decimal("156.30")
    assert abs(rev.value("SHOES") - shoes_rev) < CENT
    # ouverture 330 000 - cloture 380 000 ; achats intragroupe elimines
    assert abs(cost.value("SHOES") - Decimal(50_000)) < CENT
    assert rev.value("BESPOKE & SHOES") == rev.value("SHOES") + rev.value("BESPOKE")
    assert abs(rate.value("SHOES") - (shoes_rev + 50_000) / shoes_rev) < Decimal("1e-9")


# ------------------------------------------------------------ classification
def test_pcg_class_classifies_subsidiary_receivables(cfg):
    assert cfg.statement_of("Account Receivable - subsidiary Mage Japan") is Statement.BALANCE_SHEET
    assert cfg.statement_of("Mage Japan intercompany account") is Statement.BALANCE_SHEET


def test_unclassified_group_account_is_excluded_not_guessed(cfg, tmp_path):
    files = ms.build(tmp_path, jp_tb=ms.rebalance(list(ms.MAGE_JAPON_TB) + [("Mystery", 1000.0)]))
    wbs = read_sources(cfg, files)
    jp = next(wb for wb in wbs if wb.entity_code == "MAGE_JAPON_KK")
    jp.account_map[norm("Mystery")] = "Compte groupe inexistant"
    res = run(cfg, files, workbooks=wbs).result
    diag = [d for d in res.diagnostics if d.code == "MAP-STATEMENT-UNKNOWN"]
    assert diag and diag[0].severity is Severity.ERROR
    assert not any(ln.group_coa == "Compte groupe inexistant" for ln in res.lines)


# ---------------------------------------------------------------------- taux
def test_rates_file_csv_french_format(tmp_path):
    p = tmp_path / "taux.csv"
    p.write_text("devise;cloture;moyen;historique\nJPY;156,33;151,9425;108,0988\n"
                 "GBP;0,8391;0,8365;\n", encoding="utf-8")
    rates = load_rates_file(p)
    assert rates["JPY"] == {"eom": Decimal("156.33"), "average": Decimal("151.9425"),
                            "historical": Decimal("108.0988")}
    assert "historical" not in rates["GBP"]


def test_rates_template_roundtrip(cfg, sample_files, tmp_path):
    wbs = read_sources(cfg, sample_files)
    needs = rate_needs(cfg, wbs, "2025-12-31")
    assert [n.currency for n in needs] == ["JPY"]
    path = write_rates_template(tmp_path / "taux.xlsx", needs, "2025-12-31")
    assert load_rates_file(path)["JPY"]["eom"] == Decimal("156.33")


def test_rates_file_drives_the_run(cfg, sample_files, tmp_path):
    p = tmp_path / "taux.csv"
    p.write_text("devise;cloture;moyen\nJPY;160;158\n", encoding="utf-8")
    res = run(cfg, sample_files, RunOptions(rates_file=p)).result
    jp_ar = res.total("Account Receivable", entity="MAGE_JAPON_KK")
    assert abs(jp_ar - Decimal(15_400_000) / 160) < CENT


def test_apply_rates_does_not_mutate_config(cfg):
    before = copy.deepcopy(cfg.fx)
    apply_rates(cfg, "2030-12-31", {"JPY": {"eom": "150"}})
    assert cfg.fx == before


def test_missing_rates_detected(cfg, sample_files):
    wbs = read_sources(cfg, sample_files)
    assert missing_rates(cfg, wbs, "2031-12-31") == [("JPY", ["eom", "average"])]


@pytest.mark.parametrize("value,code", [
    ("0.0064", "FX-RATE-INVERTED"),   # 1/156,33 : cotation inversee
    ("15.633", "FX-RATE-SUSPECT"),    # decimale deplacee
    ("-1", "FX-RATE-INVALID"),
])
def test_suspicious_rates_are_flagged(cfg, value, code):
    diag = check_rate(cfg, "JPY", "eom", Decimal(value), "2026-12-31")
    assert diag is not None and diag.code == code


def test_plausible_rate_passes(cfg):
    assert check_rate(cfg, "JPY", "eom", Decimal("161.2"), "2026-12-31") is None


# ------------------------------------------------------------- rapprochement
def _reference(path: Path, out, tweaks: dict[str, Decimal]) -> Path:
    """Classeur au format EC+ : libelle en B, montant en C, feuilles de detail
    en 'very hidden' comme dans le fichier reel."""
    wb = Workbook()
    summary = wb.active  # feuille de synthese visible, comme dans EC+/2025.xlsx
    summary.title = "Consolidated BS"
    summary.append([None, "Cash and cash equivalents", 999999])  # ne doit PAS etre lue
    for name, key in (("Detailed Consolidated BS", "Detailed BS"),
                      ("Detailed Consolidated PL", "Detailed PL")):
        ws = wb.create_sheet(name)
        ws.sheet_state = "veryHidden"
        ws.append(["MAGE SAS"])
        for row in out.reports[key].rows:
            if row.kind == "section":
                ws.append([None, row.caption])
                continue
            value = row.value() + tweaks.get(row.caption, 0)
            ws.append([None, row.caption, float(value)])
    wb.save(path)
    return path


def test_reconciliation_reads_very_hidden_sheets(out, tmp_path):
    ref = load_reference(_reference(tmp_path / "ref.xlsx", out, {}))
    assert norm("Cash Bank") in ref[Statement.BALANCE_SHEET]
    assert norm("Shoe Sales") in ref[Statement.PROFIT_AND_LOSS]


def test_reconciliation_identical_reference_passes(cfg, sample_files, out, tmp_path):
    ref = _reference(tmp_path / "ref.xlsx", out, {})
    res = run(cfg, sample_files, RunOptions(reference=ref))
    assert res.reconciliation.ok
    c9 = next(c for c in res.result.controls if c.code == "C9")
    assert c9.passed


def test_reconciliation_pinpoints_gaps(cfg, sample_files, out, tmp_path):
    ref = _reference(tmp_path / "ref.xlsx", out,
                     {"Cash Bank": Decimal(500), "Rent": Decimal("-0.40")})
    res = run(cfg, sample_files, RunOptions(reference=ref))
    gaps = {ln.caption: ln.difference for ln in res.reconciliation.gaps}
    # seule la ligne est alteree dans la reference (pas ses totaux) ; l'ecart de
    # 0,40 EUR sur Rent reste sous la tolerance de 1 EUR
    assert set(gaps) == {"Cash Bank"}
    assert abs(gaps["Cash Bank"] + 500) < CENT  # valeur relue en flottant
    c9 = next(c for c in res.result.controls if c.code == "C9")
    assert not c9.passed and "500,00" in c9.detail


# --------------------------------------------------------------------- export
def test_export_sheets_and_exact_values(cfg, sample_files, tmp_path):
    res = run(cfg, sample_files, RunOptions(presentation={"scale": "thousands", "decimals": 1}))
    path = export(res, tmp_path / "etats.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames[0] == "Synthèse"
    assert {"Consolidated BS", "BS by entity", "Gross margin", "Controls",
            "Audit trail"} <= set(wb.sheetnames)
    ws = wb["Consolidated BS"]
    cash = next(r for r in ws.iter_rows(min_row=7)
                if r[0].value == "Cash and cash equivalents")[1]
    # valeur EXACTE malgre l'affichage en milliers : l'echelle est dans le format
    assert abs(cash.value - float(res.bs.total("Cash and cash equivalents"))) < 0.001
    assert ",\"k" in cash.number_format


def test_export_mapping_sheet_when_unmapped(cfg, tmp_path):
    files = ms.build(tmp_path / "src", sas_tb=[("Loyer entrepot" if c == "Rent" else c, v)
                                              for c, v in ms.MAGE_SAS_TB])
    res = run(cfg, files)
    wb = load_workbook(export(res, tmp_path / "etats.xlsx"))
    ws = wb["Mapping à compléter"]
    assert ws["B5"].value == "Loyer entrepot" and ws["E5"].value == "Rent"
    assert "NE PAS DIFFUSER" in wb["Synthèse"]["A5"].value


# ------------------------------------------------------------ parametrage
def test_shipped_configuration_is_consistent(cfg):
    errors = [f for f in check_config(cfg) if f.severity is Severity.ERROR]
    assert not errors, [f.message for f in errors]


def test_check_detects_orphan_summary_reference(cfg):
    broken = copy.deepcopy(cfg)
    broken.bs.summary_sections[0]["lines"][0]["from"].append("Cash Bnak")
    codes = {f.code for f in check_config(broken)}
    assert "STRUCT-ORPHAN" in codes


def test_c10_flags_amount_outside_statements(cfg, tmp_path):
    """Une creance sur filiale non eliminee (contrepartie absente) ne figure
    sur aucune ligne des etats : C10 doit le signaler."""
    jp = [(c, v) for c, v in ms.MAGE_JAPON_TB if not c.startswith("Accounts Payable - Inter")]
    files = ms.build(tmp_path, jp_tb=ms.rebalance(jp))
    res = run(cfg, files).result
    c10 = next(c for c in res.controls if c.code == "C10")
    assert not c10.passed and "Account Receivable - subsidiary Mage Japan" in c10.detail
