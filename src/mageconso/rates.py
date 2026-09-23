"""Gestion des taux de change d'une cloture.

Trois besoins d'ergonomie, que l'edition directe de ``config/fx.yaml`` ne
couvrait pas :

1. saisir les taux dans un FICHIER simple (CSV ou Excel) plutot qu'en YAML ;
2. savoir AVANT de consolider quelles devises manquent, au lieu de decouvrir
   des lignes non converties apres coup ;
3. detecter les erreurs de saisie classiques : taux inverse (0,0064 au lieu de
   156,33 pour le yen, les classeurs Mage etant en cotation INDIRECTE
   1 EUR = X devise) ou decimale deplacee (15,633).
"""
from __future__ import annotations

import copy
import csv
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable

from .config import AppConfig
from .importers.excel import to_decimal
from .models import Diagnostic, Severity

RATE_KINDS = ("eom", "average", "historical")

# alias acceptes dans les en-tetes de fichiers de taux
_HEADERS = {
    "currency": {"currency", "devise", "ccy", "monnaie"},
    "eom": {"eom", "cloture", "clôture", "closing", "fin de mois", "taux cloture"},
    "average": {"average", "moyen", "moyenne", "avg", "taux moyen"},
    "historical": {"historical", "historique", "historic", "taux historique"},
}


@dataclass
class RateNeed:
    currency: str
    kinds: tuple[str, ...]
    entities: list[str]
    known: dict[str, Decimal | None]
    suggested: dict[str, Decimal | None]

    @property
    def missing(self) -> list[str]:
        return [k for k in self.kinds if self.known.get(k) is None]


def needed_currencies(cfg: AppConfig, workbooks: Iterable) -> dict[str, list[str]]:
    """Devises a convertir -> entites concernees (hors devise de reporting)."""
    out: dict[str, list[str]] = {}
    for wb in workbooks:
        ent = cfg.entity(wb.entity_code) if wb.entity_code else None
        cur = (wb.currency or (ent.currency if ent else None) or "").upper()
        if cur and cur != cfg.reporting_currency:
            out.setdefault(cur, []).append(wb.entity_code or wb.entity_name or wb.path.stem)
    return out


def _closings(cfg: AppConfig) -> list[str]:
    return sorted((cfg.fx.get("rates") or {}).keys())


def previous_rate(cfg: AppConfig, currency: str, kind: str,
                  before: str | None) -> tuple[str, Decimal] | None:
    """Dernier taux connu pour une devise, anterieur a ``before``."""
    for closing in reversed(_closings(cfg)):
        if before and closing >= before:
            continue
        val = ((cfg.fx.get("rates") or {}).get(closing, {}).get(kind) or {}).get(currency)
        if val is not None:
            return closing, Decimal(str(val))
    return None


def rate_needs(cfg: AppConfig, workbooks: Iterable, closing: str) -> list[RateNeed]:
    """Pour chaque devise : taux deja connus pour la cloture, et suggestion
    (dernier taux connu) pour pre-remplir une saisie."""
    needs = []
    for cur, ents in sorted(needed_currencies(cfg, workbooks).items()):
        kinds = ("eom", "average", "historical")
        known = {}
        suggested = {}
        for k in kinds:
            if k == "historical":
                v = (cfg.fx.get("historical") or {}).get(cur)
                known[k] = Decimal(str(v)) if v is not None else None
                suggested[k] = known[k]
                continue
            v = ((cfg.fx.get("rates") or {}).get(closing, {}).get(k) or {}).get(cur)
            known[k] = Decimal(str(v)) if v is not None else None
            prev = previous_rate(cfg, cur, k, closing)
            suggested[k] = known[k] if known[k] is not None else (prev[1] if prev else None)
        needs.append(RateNeed(cur, kinds, ents, known, suggested))
    return needs


def check_rate(cfg: AppConfig, currency: str, kind: str, value: Decimal,
               closing: str) -> Diagnostic | None:
    """Signale un taux vraisemblablement faux par comparaison au dernier connu."""
    if value <= 0:
        return Diagnostic("FX-RATE-INVALID", Severity.ERROR,
                          f"Taux {kind} {currency} = {value} : doit etre strictement positif.")
    ref = previous_rate(cfg, currency, kind, closing)
    if ref is None and kind == "historical":
        return None
    if ref is None:
        return None
    ref_closing, ref_value = ref
    ratio = value / ref_value
    if abs(value * ref_value - 1) < Decimal("0.35") and ref_value > 2:
        return Diagnostic(
            "FX-RATE-INVERTED", Severity.ERROR,
            f"Taux {kind} {currency} = {value} : semble INVERSE (dernier connu "
            f"{ref_value} au {ref_closing}). Les taux sont en cotation indirecte : "
            f"1 EUR = X {currency}.")
    if ratio < Decimal("0.7") or ratio > Decimal("1.43"):
        return Diagnostic(
            "FX-RATE-SUSPECT", Severity.WARNING,
            f"Taux {kind} {currency} = {value} : ecart de {((ratio - 1) * 100):+.0f} % "
            f"avec le dernier connu ({ref_value} au {ref_closing}). Verifier la "
            "saisie (decimale deplacee ?).")
    return None


def load_rates_file(path: str | Path) -> dict[str, dict[str, Decimal]]:
    """Lit un fichier de taux (CSV ; ou , , ou Excel), une ligne par devise :

        devise ; cloture ; moyen ; historique
        JPY    ; 156,33  ; 151,94 ; 108,0988

    Les en-tetes francais ou anglais sont acceptes, ainsi que la virgule
    decimale. Retourne {devise: {"eom": .., "average": .., "historical": ..}}.
    """
    p = Path(path)
    rows: list[list[object]]
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(p, data_only=True, read_only=True)
        ws = wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
    else:
        text = p.read_text(encoding="utf-8-sig")
        dialect = csv.Sniffer().sniff(text.splitlines()[0], delimiters=";,\t")
        rows = list(csv.reader(text.splitlines(), dialect))

    header_idx, cols = None, {}
    for i, row in enumerate(rows[:10]):
        found = {}
        for j, cell in enumerate(row):
            lab = str(cell or "").strip().lower()
            for key, aliases in _HEADERS.items():
                if lab in aliases:
                    found[key] = j
        if "currency" in found and len(found) > 1:
            header_idx, cols = i, found
            break
    if header_idx is None:
        raise ValueError(
            f"{p.name} : en-tetes introuvables. Attendu : devise ; cloture ; moyen "
            "; historique (ou currency ; eom ; average ; historical).")

    out: dict[str, dict[str, Decimal]] = {}
    for row in rows[header_idx + 1:]:
        if not row or cols["currency"] >= len(row):
            continue
        cur = str(row[cols["currency"]] or "").strip().upper()
        if not cur:
            continue
        entry = {}
        for kind in RATE_KINDS:
            j = cols.get(kind)
            if j is None or j >= len(row):
                continue
            val = to_decimal(row[j])
            if val is not None:
                entry[kind] = val
        if entry:
            out[cur] = entry
    return out


def apply_rates(cfg: AppConfig, closing: str,
                rates: dict[str, dict[str, Decimal | str | float]]) -> AppConfig:
    """Retourne une COPIE de la configuration completee des taux saisis.

    La configuration d'origine n'est pas modifiee : les taux d'une execution
    n'ecrasent jamais silencieusement ceux d'une autre cloture.
    """
    new = copy.deepcopy(cfg)
    fx = new.fx
    table = fx.setdefault("rates", {}).setdefault(closing, {})
    for cur, kinds in rates.items():
        for kind, value in (kinds or {}).items():
            if value in (None, ""):
                continue
            try:
                dec = value if isinstance(value, Decimal) else Decimal(str(value).replace(",", "."))
            except InvalidOperation as exc:
                raise ValueError(f"Taux {kind} {cur} illisible : {value!r}") from exc
            if kind == "historical":
                fx.setdefault("historical", {})[cur.upper()] = dec
            else:
                table.setdefault(kind, {})[cur.upper()] = dec
    return new


def write_rates_template(path: str | Path, needs: list[RateNeed], closing: str) -> Path:
    """Fichier de taux pre-rempli, a completer puis a passer avec --rates."""
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    if out.suffix.lower() == ".xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = "Taux"
        ws.append(["devise", "cloture", "moyen", "historique", "entites",
                   f"taux au {closing} - 1 EUR = X devise"])
        for c in ws[1]:
            c.font = Font(bold=True)
        fill = PatternFill("solid", fgColor="FFF2CC")
        for n in needs:
            ws.append([n.currency,
                       *[float(n.suggested[k]) if n.suggested.get(k) is not None else None
                         for k in ("eom", "average", "historical")],
                       ", ".join(n.entities)])
            for k, col in (("eom", 2), ("average", 3), ("historical", 4)):
                if n.known.get(k) is None:
                    ws.cell(row=ws.max_row, column=col).fill = fill
        ws.column_dimensions["E"].width = 40
        ws.column_dimensions["F"].width = 44
        wb.save(out)
        return out
    with out.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(["devise", "cloture", "moyen", "historique", "entites"])
        for n in needs:
            w.writerow([n.currency,
                        *[str(n.suggested[k]).replace(".", ",")
                          if n.suggested.get(k) is not None else ""
                          for k in ("eom", "average", "historical")],
                        ", ".join(n.entities)])
    return out
