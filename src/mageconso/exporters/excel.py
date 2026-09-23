"""Export Excel des etats consolides.

Principes :
 - les cellules recoivent la VALEUR EXACTE ; le format de nombre (derive de
   config/presentation.yaml) assure l'affichage, echelle comprise. Changer la
   presentation ne modifie donc jamais un montant stocke ;
 - la premiere feuille, "Synthese", repond aux trois questions de l'utilisateur
   avant qu'il ouvre quoi que ce soit : les etats sont-ils diffusables ? que
   disent-ils ? sur quelles sources et quels taux reposent-ils ?
"""
from __future__ import annotations

from pathlib import Path
from typing import TYPE_CHECKING

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..models import Severity
from ..reports import TOTAL, StatementReport

if TYPE_CHECKING:  # pragma: no cover
    from ..pipeline import RunOutput

THIN = Side(style="thin", color="808080")
DOUBLE = Side(style="double", color="404040")
GREEN, RED, AMBER, GREY = "C6EFCE", "FFC7CE", "FFEB9C", "F2F2F2"

KIND_STYLE = {
    "section": "section",
    "total": "total",
    "grand_total": "grand_total",
    "cascade_total": "total",
    "cascade": "line",
    "percent": "line",
}


def _apply(cell, spec: dict) -> None:
    if spec.get("bold") or spec.get("font_size") or spec.get("colour"):
        cell.font = Font(bold=bool(spec.get("bold")), size=spec.get("font_size") or 11,
                         color=spec.get("colour"))
    if spec.get("fill"):
        cell.fill = PatternFill("solid", fgColor=spec["fill"])
    if spec.get("top_border") == "thin":
        cell.border = Border(top=THIN)
    elif spec.get("top_border") == "double":
        cell.border = Border(top=DOUBLE)


def _fill(cell, colour: str) -> None:
    cell.fill = PatternFill("solid", fgColor=colour)


def _header_row(ws, row: int, labels: list[str], style: dict) -> None:
    for i, lab in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=lab)
        _apply(c, style or {"bold": True, "fill": "DCE6F1"})
        c.alignment = Alignment(horizontal="center" if i > 1 else "left",
                                vertical="center", wrap_text=True)


def _print_setup(ws, title_rows: str | None = None) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    if title_rows:
        ws.print_title_rows = title_rows


# --------------------------------------------------------------------- etats
def write_statement(ws, report: StatementReport, out: "RunOutput") -> None:
    fmt = out.fmt
    pres = out.cfg.presentation or {}
    styles = pres.get("styles") or {}
    export = pres.get("export") or {}
    numfmt = fmt.excel_number_format()
    pctfmt = fmt.excel_percent_format()
    # reglage de l'execution (surcharges comprises), pas la config globale
    with_pct = bool(fmt.settings.get("percent_of_revenue", True)) and bool(report.percent_base)

    ws["A1"] = "MAGE SAS"
    _apply(ws["A1"], {"bold": True, "font_size": 13})
    ws["A2"] = report.title
    _apply(ws["A2"], {"bold": True})
    ws["A3"] = f"FOR THE PERIOD ENDED {fmt.period_end_label(report.period_end)}"
    ws["A4"] = " ".join(x for x in (report.currency, fmt.scale_note()) if x)
    ws["A4"].font = Font(italic=True, color="595959")

    # colonnes : [valeur (+ %)] par colonne du rapport
    layout: list[tuple[str, str]] = []
    for col in report.columns:
        layout.append((col, "value"))
        if with_pct:
            layout.append((col, "pct"))

    header = 6
    labels = [""] + [report.label(c) if kind == "value" else "%" for c, kind in layout]
    _header_row(ws, header, labels, styles.get("header"))

    r = header + 1
    for row in report.rows:
        style = styles.get(KIND_STYLE.get(row.kind, "line")) or {}
        label_cell = ws.cell(row=r, column=1, value=row.caption)
        _apply(label_cell, style)
        if row.level:
            label_cell.alignment = Alignment(indent=row.level)
        if row.kind != "section":
            for j, (col, kind) in enumerate(layout, start=2):
                cell = ws.cell(row=r, column=j)
                if kind == "pct":
                    pct = report.percent(row, col)
                    if pct is not None:
                        cell.value = float(pct)
                        cell.number_format = pctfmt
                        cell.font = Font(italic=True, color="595959")
                    continue
                val = row.values.get(col)
                if val is None:
                    continue
                cell.value = float(val)
                cell.number_format = pctfmt if row.kind == "percent" else numfmt
                _apply(cell, style)
                if col == TOTAL and len(report.columns) > 1:
                    cell.font = Font(bold=True)
        r += 1

    ws.column_dimensions["A"].width = export.get("column_width_label", 52)
    for j, (_, kind) in enumerate(layout, start=2):
        ws.column_dimensions[get_column_letter(j)].width = (
            8 if kind == "pct" else export.get("column_width_value", 16))
    ws.row_dimensions[header].height = 30
    if export.get("freeze_panes", True):
        ws.freeze_panes = ws.cell(row=header + 1, column=2)
    _print_setup(ws, f"{header}:{header}")


# ------------------------------------------------------------------ synthese
def write_summary(ws, out: "RunOutput") -> None:
    fmt = out.fmt
    res = out.result
    numfmt = fmt.excel_number_format()
    ws["A1"] = "MAGE SAS - CONSOLIDATION"
    _apply(ws["A1"], {"bold": True, "font_size": 15})
    ws["A2"] = f"Période close le {fmt.period_end_label(res.period.end_date)}"
    ws["A3"] = f"Produit le {out.started_at:%d/%m/%Y à %H:%M}"
    ws["A3"].font = Font(italic=True, color="595959")

    blocking = out.blocking
    status = ws["A5"]
    if blocking:
        status.value = (f"NE PAS DIFFUSER - {len(blocking)} contrôle(s) bloquant(s) "
                        "en échec. Voir l'onglet Controls.")
        _fill(status, RED)
    else:
        warn = sum(1 for c in res.controls if not c.passed)
        status.value = ("Contrôles bloquants satisfaits"
                        + (f" - {warn} avertissement(s) à examiner" if warn else ""))
        _fill(status, AMBER if warn else GREEN)
    status.font = Font(bold=True, size=12)
    ws.merge_cells("A5:F5")

    r = 7
    ws.cell(row=r, column=1, value="Chiffres clés").font = Font(bold=True, size=12)
    r += 1
    for label, value in out.kpis():
        ws.cell(row=r, column=1, value=label)
        c = ws.cell(row=r, column=2, value=float(value))
        c.number_format = numfmt
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Sources").font = Font(bold=True, size=12)
    r += 1
    _header_row(ws, r, ["Fichier", "Entité", "Devise", "Clôture", "Lignes",
                        "Balance équilibrée"], {"bold": True, "fill": "DCE6F1"})
    for wb in out.workbooks:
        r += 1
        total = sum((ln.amount for ln in wb.lines), 0)
        vals = [wb.path.name, wb.entity_code or wb.entity_name or "?", wb.currency,
                wb.closing, len(wb.lines), "oui" if abs(total) < 0.01 else f"NON ({total})"]
        for j, v in enumerate(vals, start=1):
            ws.cell(row=r, column=j, value=v)
        _fill(ws.cell(row=r, column=6), GREEN if abs(total) < 0.01 else RED)

    rates = out.rates_used()
    if rates:
        r += 2
        ws.cell(row=r, column=1, value=f"Taux appliqués (1 EUR = X devise) au {out.closing}"
                ).font = Font(bold=True, size=12)
        r += 1
        _header_row(ws, r, ["Devise", "Type", "Taux"], {"bold": True, "fill": "DCE6F1"})
        names = {"eom": "clôture", "average": "moyen", "historical": "historique"}
        for cur, kind, rate in rates:
            r += 1
            ws.cell(row=r, column=1, value=cur)
            ws.cell(row=r, column=2, value=names.get(kind, kind))
            ws.cell(row=r, column=3, value=float(rate)).number_format = "0.0000"

    r += 2
    ws.cell(row=r, column=1, value="Contrôles").font = Font(bold=True, size=12)
    for c in res.controls:
        r += 1
        ws.cell(row=r, column=1, value=f"{c.code}  {c.label}")
        s = ws.cell(row=r, column=2, value="OK" if c.passed else
                    ("ÉCHEC" if c.severity is Severity.ERROR else "À VOIR"))
        _fill(s, GREEN if c.passed else (RED if c.severity is Severity.ERROR else AMBER))
        ws.cell(row=r, column=3, value=c.detail)

    counts = out.diagnostics_by_severity()
    if counts:
        r += 2
        ws.cell(row=r, column=1, value="Diagnostics").font = Font(bold=True, size=12)
        for sev in ("ERROR", "WARNING", "INFO"):
            if counts.get(sev):
                r += 1
                ws.cell(row=r, column=1, value={"ERROR": "Erreurs", "WARNING":
                        "Avertissements", "INFO": "Informations"}[sev])
                ws.cell(row=r, column=2, value=counts[sev])

    ws.column_dimensions["A"].width = 58
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 60
    for col in "DEF":
        ws.column_dimensions[col].width = 14
    _print_setup(ws)


# ------------------------------------------------------------ rapprochement
def write_reconciliation(ws, out: "RunOutput") -> None:
    rec = out.reconciliation
    numfmt = out.fmt.excel_number_format()
    ws["A1"] = f"Rapprochement avec {rec.reference_file}"
    ws["A1"].font = Font(bold=True, size=13)
    counts = rec.counts()
    ws["A2"] = (f"{counts.get('OK', 0)} conforme(s), {len(rec.gaps)} écart(s) au-delà "
                f"de {rec.tolerance} EUR")
    _header_row(ws, 4, ["État", "Poste", "Référence", "Calculé", "Écart", "Écart %",
                        "Statut", "Source"], {"bold": True, "fill": "DCE6F1"})
    order = {"ECART": 0, "absent du calcul": 1, "absent de la reference": 2, "OK": 3}
    lines = sorted(rec.lines, key=lambda ln: (order.get(ln.status(rec.tolerance), 9),
                                              -(abs(ln.difference or 0))))
    for r, ln in enumerate(lines, start=5):
        st = ln.status(rec.tolerance)
        ws.cell(row=r, column=1, value="Bilan" if ln.statement.value.startswith("BAL") else "Résultat")
        ws.cell(row=r, column=2, value=ln.caption)
        for j, v in ((3, ln.reference), (4, ln.computed), (5, ln.difference)):
            if v is not None:
                ws.cell(row=r, column=j, value=float(v)).number_format = numfmt
        if ln.difference_pct is not None:
            ws.cell(row=r, column=6, value=float(ln.difference_pct)).number_format = "0.0%"
        s = ws.cell(row=r, column=7, value=st)
        _fill(s, GREEN if st == "OK" else (RED if st == "ECART" else AMBER))
        ws.cell(row=r, column=8, value=ln.source)
    ws.column_dimensions["B"].width = 50
    for col in "CDE":
        ws.column_dimensions[col].width = 16
    ws.column_dimensions["G"].width = 22
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:H{max(4, len(lines) + 4)}"


# ------------------------------------------------------- mapping a completer
def write_unmapped(ws, out: "RunOutput") -> None:
    ws["A1"] = "Comptes sources NON MAPPÉS - exclus du consolidé"
    ws["A1"].font = Font(bold=True, size=13, color="9C0006")
    ws["A2"] = ("Validez ou corrigez la colonne 'Compte groupe proposé', puis reportez "
                "la correspondance dans la feuille 'Mapping accounts' du fichier source "
                "ou dans config/mapping/<entite>.csv.")
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A2:G2")
    ws.row_dimensions[2].height = 32
    _header_row(ws, 4, ["Entité", "Compte local", "Montant", "Devise",
                        "Compte groupe proposé", "Autres suggestions", "Fichier"],
                {"bold": True, "fill": "DCE6F1"})
    for r, u in enumerate(out.result.unmapped, start=5):
        ws.cell(row=r, column=1, value=u.entity)
        ws.cell(row=r, column=2, value=u.local_account)
        ws.cell(row=r, column=3, value=float(u.amount)).number_format = "#,##0.00"
        ws.cell(row=r, column=4, value=u.currency)
        best = ws.cell(row=r, column=5, value=u.best or "(aucune suggestion)")
        _fill(best, AMBER)
        ws.cell(row=r, column=6, value=" | ".join(
            f"{c} ({s:.0%})" for c, s in u.suggestions[1:]))
        ws.cell(row=r, column=7, value=u.source_file)
    for col, w in zip("ABCDEFG", (18, 40, 16, 8, 40, 60, 34)):
        ws.column_dimensions[col].width = w


# ------------------------------------------------------------- techniques
def write_controls(ws, out: "RunOutput") -> None:
    numfmt = out.fmt.excel_number_format()
    _header_row(ws, 1, ["Code", "Contrôle", "Statut", "Gravité", "Attendu", "Obtenu",
                        "Écart", "Détail"], {"bold": True, "fill": "DCE6F1"})
    for r, ctl in enumerate(out.result.controls, start=2):
        ws.cell(row=r, column=1, value=ctl.code)
        ws.cell(row=r, column=2, value=ctl.label)
        status = ws.cell(row=r, column=3, value="OK" if ctl.passed else "ÉCHEC")
        _fill(status, GREEN if ctl.passed else
              (RED if ctl.severity is Severity.ERROR else AMBER))
        ws.cell(row=r, column=4, value=ctl.severity.value)
        for j, v in ((5, ctl.expected), (6, ctl.actual), (7, ctl.difference)):
            if v is not None:
                ws.cell(row=r, column=j, value=float(v)).number_format = numfmt
        ws.cell(row=r, column=8, value=ctl.detail).alignment = Alignment(wrap_text=True)
    ws.column_dimensions["B"].width = 58
    ws.column_dimensions["H"].width = 80
    for col in "EFG":
        ws.column_dimensions[col].width = 16
    ws.freeze_panes = "A2"


def write_diagnostics(ws, out: "RunOutput") -> None:
    _header_row(ws, 1, ["Gravité", "Code", "Entité", "Fichier", "Message"],
                {"bold": True, "fill": "DCE6F1"})
    rank = {"ERROR": 0, "WARNING": 1, "INFO": 2}
    diags = sorted(out.result.diagnostics, key=lambda d: rank.get(d.severity.value, 9))
    for r, d in enumerate(diags, start=2):
        s = ws.cell(row=r, column=1, value=d.severity.value)
        _fill(s, RED if d.severity is Severity.ERROR else
              (AMBER if d.severity is Severity.WARNING else GREY))
        ws.cell(row=r, column=2, value=d.code)
        ws.cell(row=r, column=3, value=d.entity)
        ws.cell(row=r, column=4, value=d.source_file)
        ws.cell(row=r, column=5, value=d.message).alignment = Alignment(wrap_text=True)
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 30
    ws.column_dimensions["E"].width = 110
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{max(1, len(diags) + 1)}"


def write_audit(ws, out: "RunOutput") -> None:
    heads = [
        "Entité", "Période", "Fichier source", "Onglet source", "Ligne",
        "Compte source", "Compte consolidé", "Centre de coûts",
        "Montant source", "Devise", "Inversion de signe", "Taux",
        "Type de taux", "Montant EUR", "Transformations", "Origine",
    ]
    _header_row(ws, 1, heads, {"bold": True, "fill": "DCE6F1"})
    for r, a in enumerate(out.result.audit, start=2):
        vals = [a.entity, a.period, a.source_file, a.source_sheet, a.source_row,
                a.local_account, a.group_coa, a.cost_centre, float(a.amount_source),
                a.currency, "oui" if a.sign_flip else "non",
                float(a.rate) if a.rate is not None else None, a.rate_type,
                float(a.amount_eur), a.transformations, a.origin]
        for j, v in enumerate(vals, start=1):
            ws.cell(row=r, column=j, value=v)
        ws.cell(row=r, column=9).number_format = "#,##0.00"
        ws.cell(row=r, column=14).number_format = "#,##0.00"
    for col, width in (("C", 34), ("F", 44), ("G", 44), ("O", 58)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:P{max(1, len(out.result.audit) + 1)}"


# --------------------------------------------------------------------- export
def export_workbook(path: str | Path, out: "RunOutput") -> Path:
    export = (out.cfg.presentation or {}).get("export") or {}
    wb = Workbook()
    write_summary(wb.active, out)
    wb.active.title = "Synthèse"

    for name, report in out.reports.items():
        write_statement(wb.create_sheet(title=name[:31]), report, out)

    if out.reconciliation is not None:
        write_reconciliation(wb.create_sheet("Rapprochement"), out)
    if out.result.unmapped:
        write_unmapped(wb.create_sheet("Mapping à compléter"), out)
    if export.get("include_controls_sheet", True):
        write_controls(wb.create_sheet("Controls"), out)
    if out.result.diagnostics:
        write_diagnostics(wb.create_sheet("Diagnostics"), out)
    if export.get("include_audit_sheet", True) and out.result.audit:
        write_audit(wb.create_sheet("Audit trail"), out)

    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)
    return target
