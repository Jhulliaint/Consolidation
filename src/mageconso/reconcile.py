"""Rapprochement des etats produits avec un classeur consolide de reference.

C'est l'outil de la premiere mise en service : on consolide une cloture deja
produite par le processus Excel, puis on compare ligne a ligne. Chaque ecart est
soit une erreur de l'application, soit une regle mal reconstituee - dans les deux
cas, il doit etre explique avant toute mise en production.

Format de reference attendu : celui de "EC+/<annee>.xlsx" - feuilles
"Detailed Consolidated BS" / "Detailed Consolidated PL" (souvent en
``very hidden``, lues quand meme), libelle en colonne B, montant en colonne C
(les totaux de section sont en colonne D).
"""
from __future__ import annotations

from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from .config import norm
from .importers.excel import to_decimal
from .models import ZERO, Statement
from .reports import StatementReport

REFERENCE_SHEETS = {
    Statement.BALANCE_SHEET: ["Detailed Consolidated BS", "Consolidated BS"],
    Statement.PROFIT_AND_LOSS: ["Detailed Consolidated PL", "Consolidated PL"],
}

# libelles de structure, sans montant a rapprocher
_SKIP = {"assets", "current assets", "non current assets", "revenue",
         "cost of revenue", "liabilities and shareholders' equity",
         "current liabilities", "non-current liabilities", "shareholders' equity",
         "selling, marketing & administrative expenses", "control", "control :",
         "for the period ended", "mage sas", "ytd", "€"}


@dataclass
class ReferenceValue:
    caption: str
    value: Decimal
    sheet: str
    row: int


@dataclass
class ReconciliationLine:
    statement: Statement
    caption: str
    reference: Decimal | None
    computed: Decimal | None
    source: str = ""

    @property
    def difference(self) -> Decimal | None:
        if self.reference is None or self.computed is None:
            return None
        return self.computed - self.reference

    @property
    def difference_pct(self) -> Decimal | None:
        d = self.difference
        if d is None or not self.reference:
            return None
        return d / abs(self.reference)

    def status(self, tolerance: Decimal) -> str:
        if self.reference is None:
            return "absent de la reference"
        if self.computed is None:
            return "absent du calcul"
        return "OK" if abs(self.difference) <= tolerance else "ECART"


@dataclass
class Reconciliation:
    reference_file: str
    tolerance: Decimal
    lines: list[ReconciliationLine] = field(default_factory=list)
    sheets_found: list[str] = field(default_factory=list)

    def counts(self) -> dict[str, int]:
        out: dict[str, int] = {}
        for ln in self.lines:
            s = ln.status(self.tolerance)
            out[s] = out.get(s, 0) + 1
        return out

    @property
    def gaps(self) -> list[ReconciliationLine]:
        return [ln for ln in self.lines if ln.status(self.tolerance) == "ECART"]

    @property
    def max_gap(self) -> Decimal:
        return max((abs(ln.difference) for ln in self.gaps), default=ZERO)

    @property
    def ok(self) -> bool:
        return not self.gaps


def load_reference(path: str | Path) -> dict[Statement, dict[str, ReferenceValue]]:
    """Extrait {etat: {libelle normalise: valeur}} d'un classeur de reference.

    Pour chaque etat, la feuille de DETAIL est preferee a la feuille de synthese.
    Sur chaque ligne, le libelle est la premiere cellule texte des colonnes A a C,
    le montant la premiere cellule numerique a sa droite.
    """
    wb = load_workbook(Path(path), data_only=True, read_only=True)
    out: dict[Statement, dict[str, ReferenceValue]] = {}
    titles = {norm(ws.title): ws for ws in wb.worksheets}
    for statement, candidates in REFERENCE_SHEETS.items():
        ws = next((titles[norm(c)] for c in candidates if norm(c) in titles), None)
        if ws is None:
            continue
        values: dict[str, ReferenceValue] = {}
        for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            caption, col = None, None
            for j, cell in enumerate(row[:3]):
                if isinstance(cell, str) and cell.strip():
                    caption, col = cell.strip(), j
                    break
            if caption is None or norm(caption) in _SKIP:
                continue
            value = None
            for cell in row[col + 1: col + 4]:
                if isinstance(cell, (int, float, Decimal)) and not isinstance(cell, bool):
                    value = Decimal(str(cell))
                    break
                if isinstance(cell, str) and to_decimal(cell) is not None:
                    value = to_decimal(cell)
                    break
            if value is None:
                continue
            values.setdefault(norm(caption), ReferenceValue(caption, value, ws.title, r_idx))
        out[statement] = values
    wb.close()
    return out


def reconcile(reference_path: str | Path,
              reports: dict[Statement, StatementReport],
              tolerance: Decimal = Decimal("1.00")) -> Reconciliation:
    """Compare les etats DETAILLES produits a la reference.

    Les montants compares sont les montants PRESENTES (signes de section
    appliques), comme dans le classeur de reference.
    """
    ref = load_reference(reference_path)
    rec = Reconciliation(reference_file=Path(reference_path).name, tolerance=tolerance)
    for statement, report in reports.items():
        values = ref.get(statement)
        if values is None:
            continue
        rec.sheets_found.append(statement.value)
        computed: dict[str, tuple[str, Decimal]] = {}
        restated: set[str] = set()
        for row in report.rows:
            if row.kind == "section":
                continue
            computed.setdefault(norm(row.caption), (row.caption, row.value()))
            if row.kind.startswith("cascade"):
                # reprise d'un total dans la cascade (Revenue = Total revenue) :
                # comparee si la reference la porte, jamais signalee absente
                restated.add(norm(row.caption))
        seen = set()
        for key, rv in values.items():
            seen.add(key)
            cap, val = computed.get(key, (rv.caption, None))
            rec.lines.append(ReconciliationLine(
                statement, rv.caption, rv.value, val,
                source=f"{rv.sheet}!L{rv.row}"))
        for key, (cap, val) in computed.items():
            if key not in seen and key not in restated and val != ZERO:
                rec.lines.append(ReconciliationLine(statement, cap, None, val))
    return rec
